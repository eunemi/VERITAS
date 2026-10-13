"""Offline regressions for bounded document and image intake."""

from __future__ import annotations

import asyncio
import io
import time
import zipfile
from dataclasses import dataclass
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from httpx import AsyncClient
from pypdf import PdfWriter

from app.api.v1.routes import files
from app.core.config import Settings, get_settings
from app.core.errors import ConfigurationError, PayloadTooLargeError, ValidationError
from app.desks.image import ImageDesk
from app.domain import Artifact, ArtifactKind
from app.graph.verdict import Judgement, Ruling
from app.vision import context

V1 = "/api/v1"


def docx_with_ordered_blocks() -> bytes:
    document = Document()
    document.add_paragraph("Before the table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    document.add_paragraph("After the table")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header evidence"
    section.footer.paragraphs[0].text = "Footer evidence"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(*, pages: int = 1, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    if encrypted:
        writer.encrypt("test-password")
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def docx_with_linked_sections() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Repeated header"
    document.add_section(WD_SECTION.NEW_PAGE)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def docx_with_merged_and_nested_cells() -> bytes:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged evidence"
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested evidence"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


async def upload_document(
    client: AsyncClient, filename: str, content: bytes
) -> object:
    return await client.post(
        f"{V1}/files/extract-text",
        files={"file": (filename, content, "application/octet-stream")},
    )


async def test_docx_keeps_paragraph_table_order_and_headers(
    client: AsyncClient,
) -> None:
    response = await upload_document(client, "ordered.docx", docx_with_ordered_blocks())

    assert response.status_code == 200
    text = response.json()["text"]
    assert text.index("Before the table") < text.index("First cell | Second cell")
    assert text.index("First cell | Second cell") < text.index("After the table")
    assert text.index("After the table") < text.index("Header evidence")
    assert text.index("Header evidence") < text.index("Footer evidence")


async def test_docx_does_not_repeat_linked_headers(client: AsyncClient) -> None:
    response = await upload_document(
        client, "linked-sections.docx", docx_with_linked_sections()
    )

    assert response.status_code == 200
    assert response.json()["text"].count("Repeated header") == 1


async def test_docx_does_not_repeat_merged_cells_and_keeps_nested_tables(
    client: AsyncClient,
) -> None:
    response = await upload_document(
        client, "nested-table.docx", docx_with_merged_and_nested_cells()
    )

    assert response.status_code == 200
    text = response.json()["text"]
    assert text.count("Merged evidence") == 1
    assert "Nested evidence" in text


def test_docx_zip_expansion_is_rejected_before_extraction() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 128)

    with pytest.raises(PayloadTooLargeError, match="expand"):
        files._check_docx_archive(stream.getvalue(), max_uncompressed_bytes=64)


async def test_text_cap_is_applied_incrementally(
    client, app, settings, monkeypatch
) -> None:
    small = settings.model_copy(update={"MAX_TEXT_CHARS": 20})
    monkeypatch.setitem(app.dependency_overrides, get_settings, lambda: small)

    response = await upload_document(client, "large.txt", b"x" * 21)

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "payload_too_large"


async def test_scanned_encrypted_and_malformed_pdfs_are_honest_errors(
    client: AsyncClient,
) -> None:
    scanned = await upload_document(client, "scan.pdf", pdf_bytes())
    encrypted = await upload_document(
        client, "locked.pdf", pdf_bytes(encrypted=True)
    )
    malformed = await upload_document(client, "broken.pdf", b"%PDF-1.7\nnot a PDF")

    assert scanned.status_code == 422
    assert "OCR" in scanned.json()["error"]["message"]
    assert encrypted.status_code == 422
    assert "Encrypted" in encrypted.json()["error"]["message"]
    assert malformed.status_code == 400


def test_pdf_page_cap_stops_before_extracting_the_tail(monkeypatch) -> None:
    original = files.MAX_PDF_PAGES
    files.MAX_PDF_PAGES = 1
    try:
        with pytest.raises(PayloadTooLargeError, match="too many pages"):
            files._extract_document(
                pdf_bytes(pages=2), ".pdf", max_text_chars=1000, max_upload_bytes=10000
            )
    finally:
        files.MAX_PDF_PAGES = original


def test_pdf_content_stream_cap_is_checked_before_extraction() -> None:
    class Stream:
        def get(self, key: str, default: object) -> object:
            return 128 if key == "/Length" else default

        def get_object(self) -> Stream:
            return self

    page = SimpleNamespace(get=lambda _key: Stream())

    with pytest.raises(PayloadTooLargeError, match="content"):
        files._check_pdf_content_stream(page, limit=64)


async def test_document_parsing_does_not_block_the_event_loop(
    client: AsyncClient, monkeypatch
) -> None:
    def slow_parser(*args: object, **kwargs: object) -> str:
        time.sleep(0.08)
        return "parsed in a worker"

    monkeypatch.setattr(files, "_extract_document", slow_parser)
    request = asyncio.create_task(upload_document(client, "slow.pdf", b"bytes"))
    ticks = 0

    while not request.done():
        ticks += 1
        await asyncio.sleep(0.005)

    response = await request
    assert response.status_code == 200
    assert ticks >= 3


def test_image_header_is_checked_before_image_validation(monkeypatch) -> None:
    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(stream, "PNG")
    monkeypatch.setattr(context, "MAX_PIXELS", 100)

    with pytest.raises(ValidationError, match="25 megapixels"):
        context.validate(stream.getvalue())


def test_prepared_dimensions_match_the_normalized_jpeg() -> None:
    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", (2000, 1000), "white").save(stream, "PNG")

    prepared = context.prepare(stream.getvalue())

    assert (prepared.width, prepared.height) == (1800, 900)


async def test_external_media_url_is_not_treated_as_a_local_upload(
    settings: Settings, monkeypatch
) -> None:
    remote = AsyncMock(return_value=b"remote image")
    monkeypatch.setattr("app.desks.image.fetch.fetch", remote)
    desk = ImageDesk(settings=settings)

    data = await desk._fetch(
        Artifact(
            kind=ArtifactKind.IMAGE,
            url="https://external.example/api/v1/files/media/"
            "0123456789abcdef0123456789abcdef.png",
        )
    )

    assert data == b"remote image"
    remote.assert_awaited_once()


@dataclass
class _CaptionGraph:
    seen: list[Artifact]

    async def run(self, artifact: Artifact) -> object:
        self.seen.append(artifact)
        return SimpleNamespace(
            ruling=Ruling(
                judgement=Judgement.UNCERTAIN,
                confidence=0,
                headline="Caption checked",
                rationale="The caption is not independently established.",
            ),
            reasoning=(),
            state={},
        )


class _UnavailableReader:
    async def read(self, image: bytes) -> object:
        raise ConfigurationError("OCR is unavailable")


async def test_caption_is_checked_when_ocr_is_unavailable(
    monkeypatch, settings
) -> None:
    graph = _CaptionGraph([])
    desk = ImageDesk(settings=settings, graph=graph)
    monkeypatch.setattr(desk, "_fetch", AsyncMock(return_value=b"image"))
    monkeypatch.setattr(
        "app.desks.image.get_image_reader", lambda _settings: _UnavailableReader()
    )

    report = await desk.examine(
        Artifact(
            kind=ArtifactKind.IMAGE,
            url="https://example.test/image.png",
            content="The bridge opened in March 2026.",
        )
    )

    assert graph.seen[0].content == "The bridge opened in March 2026."
    assert "unreadable" in report.verdict.rationale.lower()
    assert report.detail is not None
    assert "unreadable" in report.detail.limitations[0].lower()


async def test_vision_trigger_uses_key_for_vision_base_url(
    monkeypatch, settings
) -> None:
    settings = settings.model_copy(
        update={
            "VISION_BASE_URL": "https://api.groq.com/openai/v1",
            "GROQ_API_KEY": "groq-test-key",
        }
    )
    desk = ImageDesk(settings=settings)
    inspect_image = AsyncMock(return_value=object())
    monkeypatch.setattr(desk, "_fetch", AsyncMock(return_value=b"image"))
    monkeypatch.setattr(desk, "_inspect", inspect_image)

    result = await desk.examine(
        Artifact(kind=ArtifactKind.IMAGE, url="https://example.test/image.png")
    )

    assert result is inspect_image.return_value
    inspect_image.assert_awaited_once()
